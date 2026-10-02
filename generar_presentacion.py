#!/usr/bin/env python3
"""Genera la presentación (10 diapositivas) a partir del Excel del estudio de mercado.

Uso:  python generar_presentacion.py [estudio.xlsx] [salida.pptx]

Lee los VALORES CALCULADOS guardados por Excel (guarde el libro antes de ejecutar).
Los colores replican el libro: azul marino 1F3864, azul 2F5597, grises BFBFBF/D9D9D9,
verde claro E2EFDA (resultado clave) y amarillo FFF2CC (entrada).
"""
import sys
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from lxml import etree
from matplotlib.patches import Wedge

import openpyxl
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ---------- Identidad ----------
AVALUADOR = "Andrés Felipe Cuartas Montoya"
RAA = "RAA 1.128.391.782"

NAVY, BLUE = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x2F, 0x55, 0x97)
G1, G2 = RGBColor(0xBF, 0xBF, 0xBF), RGBColor(0xD9, 0xD9, 0xD9)
GREEN, YELLOW = RGBColor(0xE2, 0xEF, 0xDA), RGBColor(0xFF, 0xF2, 0xCC)
OKGREEN, RED = RGBColor(0x37, 0x86, 0x3C), RGBColor(0xC0, 0x39, 0x2B)
WHITE, INK, MUTED = RGBColor(255, 255, 255), RGBColor(0x26, 0x2B, 0x33), RGBColor(0x6B, 0x72, 0x80)
FONT = "Calibri"
W, H = Inches(13.333), Inches(7.5)


# ---------- Utilidades de formato ----------
def num(v, d=0):
    return f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def cop(v):
    return "$ " + num(v)


def mill(v, d=1):
    return f"$ {num(v / 1e6, d)} M"


def pct(v, d=1):
    return num(v * 100, d) + " %"


def fnum(v, default=0.0):
    return float(v) if isinstance(v, (int, float)) else default


# ---------- Lectura del Excel ----------
def leer(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    m, c, d, l, p = (wb[n] for n in ("MERCADO RURAL", "COMBINACIONES", "DEPRECIACIÓN", "LIQUIDACIÓN", "PROYECCIÓN"))
    D = {}
    D["fecha"] = "/".join(str(m[a].value) for a in ("M2", "L2", "K2") if m[a].value is not None) if False else \
        f"{int(fnum(m['M2'].value)):02d}/{int(fnum(m['L2'].value)):02d}/{int(fnum(m['K2'].value))}"
    D["solicitante"] = m["F3"].value or "—"
    D["direccion"] = m["E4"].value or "—"
    D["depto"], D["municipio"] = m["D6"].value or "—", m["H6"].value or "—"
    D["referencia"] = str(d["A2"].value or "").split("·")[-1].strip()

    ofertas, r = [], 9
    while isinstance(m.cell(r, 1).value, (int, float)):
        g = lambda col: m.cell(r, col).value
        ofertas.append(dict(no=int(g(1)), ubic=g(2), pedido=fnum(g(4)), neg=fnum(g(5)), depurado=fnum(g(6)),
                            area=fnum(g(7)), ha=fnum(g(8)), constr=fnum(g(9)), desc=g(10), fuente=g(11),
                            fecha=g(15), comp=str(g(16) or "")))
        r += 1
    D["ofertas"] = ofertas
    D["obs"] = next((str(m.cell(x, 1).value) for x in range(r, r + 40)
                     if str(m.cell(x, 1).value or "").startswith("OBSERVACIONES")), "")

    est = {}
    for x in range(r, r + 12):
        k = str(m.cell(x, 6).value or "")
        if k:
            est[k.split(" (")[0]] = m.cell(x, 8).value
    D["est"] = est

    D["cv_max"] = fnum(c["B5"].value, 0.10)
    D["filtro"] = c["B6"].value
    D["elegibles"] = c["B7"].value
    D["combos"] = [dict(nombre=c.cell(x, 1).value, n=c.cell(x, 2).value, ofertas=c.cell(x, 3).value,
                        prom=fnum(c.cell(x, 6).value), de=fnum(c.cell(x, 7).value), cv=fnum(c.cell(x, 8).value),
                        cumple=c.cell(x, 12).value, m2=fnum(c.cell(x, 14).value))
                   for x in range(13, 16) if c.cell(x, 1).value]
    D["sel"] = dict(nombre=c["B18"].value, ofertas=c["B19"].value, cv=fnum(c["B20"].value), cumple=c["B21"].value,
                    ha=fnum(c["B22"].value), m2=fnum(c["B23"].value))

    D["proy"] = dict(baja=fnum(p["D5"].value), alta=fnum(p["F5"].value), amp=fnum(p["H5"].value),
                     estado=p["L5"].value, resumen=p["L8"].value, nota=p["A21"].value,
                     nbaja=p["D8"].value, nalta=p["F8"].value,
                     ref=[[p.cell(x, k).value for k in range(1, 8)] for x in (18, 19)])

    constr = []
    for x in range(5, 15):
        if d.cell(x, 2).value:
            constr.append(dict(item=d.cell(x, 2).value, edad=d.cell(x, 3).value, vida=d.cell(x, 4).value,
                               estado=d.cell(x, 5).value, repos=fnum(d.cell(x, 6).value), fd=fnum(d.cell(x, 12).value),
                               adop=fnum(d.cell(x, 14).value), area=fnum(d.cell(x, 15).value),
                               total=fnum(d.cell(x, 16).value), oferta=d.cell(x, 1).value))
    D["constr"] = constr
    D["constr_total"] = fnum(d["P15"].value)

    liq = dict(terreno=None, principal=[], anexas=[], tot_p=0, tot_a=0, total=0, letras="")
    sec = None
    for x in range(3, 30):
        b = str(l.cell(x, 2).value or "")
        if b.startswith("Notas"):
            break
        row = (b, l.cell(x, 3).value, fnum(l.cell(x, 4).value), fnum(l.cell(x, 5).value), fnum(l.cell(x, 6).value))
        if b == "TERRENO":
            sec = "t"
        elif b == "CONSTRUCCIÓN PRINCIPAL":
            sec = "p"
        elif b == "CONSTRUCCIONES ANEXAS":
            sec = "a"
        elif b == "DESCRIPCIÓN":
            continue
        elif b.startswith("VALOR TOTAL DE LA CONSTRUCCIÓN PRINCIPAL"):
            liq["tot_p"] = row[4]
        elif b.startswith("VALOR TOTAL DE LAS CONSTRUCCIONES ANEXAS"):
            liq["tot_a"] = row[4]
        elif b == "VALOR TOTAL DEL AVALÚO":
            liq["total"] = row[4]
        elif b == "VALOR TOTAL EN LETRAS":
            liq["letras"] = str(l.cell(x, 3).value or "")
        elif b == "TERRENO":
            continue
        elif b and sec == "t" and liq["terreno"] is None:
            liq["terreno"] = row
        elif b and sec == "p":
            liq["principal"].append(row)
        elif b and sec == "a":
            liq["anexas"].append(row)
    D["liq"] = liq
    return D


# ---------- Paleta ampliada ----------
BG = RGBColor(0xF3, 0xF5, 0xF9)
INK2 = RGBColor(0x3A, 0x44, 0x57)
LINE = RGBColor(0xE1, 0xE6, 0xEF)
ACCENT = RGBColor(0x92, 0xD0, 0x50)
DEEP = RGBColor(0x0F, 0x1F, 0x3D)
AMBER = RGBColor(0xE0, 0xA1, 0x00)
SKY = RGBColor(0x8F, 0xAA, 0xDC)
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


# ---------- Imágenes generadas (fondo topográfico y medidor) ----------
def hexc(c):
    return "#%02X%02X%02X" % (c[0], c[1], c[2])


def img_topo(path):
    rng = np.random.default_rng(7)
    x, y = np.meshgrid(np.linspace(0, 13.333, 420), np.linspace(0, 7.5, 240))
    z = 0.12 * x + 0.05 * y
    for _ in range(9):
        cx, cy = rng.uniform(0, 13.3), rng.uniform(0, 7.5)
        sx, sy = rng.uniform(1.2, 3.2), rng.uniform(0.9, 2.4)
        z += rng.uniform(0.5, 1.4) * np.exp(-(((x - cx) / sx) ** 2 + ((y - cy) / sy) ** 2))
    fig = plt.figure(figsize=(13.333, 7.5), dpi=110)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.contour(x, y, z, levels=26, colors=[hexc(SKY)], linewidths=0.8, alpha=0.30)
    ax.contour(x, y, z, levels=6, colors=[hexc(ACCENT)], linewidths=1.3, alpha=0.30)
    ax.set_xlim(0, 13.333)
    ax.set_ylim(7.5, 0)
    ax.axis("off")
    fig.savefig(path, transparent=True)
    plt.close(fig)


def img_gauge(path, cv, limit):
    top = max(cv * 1.25, limit * 3)
    fig = plt.figure(figsize=(4.2, 2.35), dpi=200)
    ax = fig.add_axes([0.02, 0.02, 0.96, 0.96])
    zonas = [(0, limit, OKGREEN), (limit, 2 * limit, AMBER), (2 * limit, top, RED)]
    for a, b, col in zonas:
        a, b = min(a, top), min(b, top)
        if b > a:
            ax.add_patch(Wedge((0, 0), 1.0, 180 - 180 * b / top, 180 - 180 * a / top, width=0.30,
                               facecolor=hexc(col), edgecolor="white", linewidth=2))
    ang = np.pi * (1 - min(cv, top) / top)
    ax.plot([0, 0.80 * np.cos(ang)], [0, 0.80 * np.sin(ang)], color=hexc(NAVY), lw=4, solid_capstyle="round")
    ax.add_patch(plt.Circle((0, 0), 0.09, color=hexc(NAVY)))
    ax.text(-0.88, -0.14, "0 %", ha="center", fontsize=9, color=hexc(MUTED))
    ax.text(0.88, -0.14, f"{top * 100:.0f} %", ha="center", fontsize=9, color=hexc(MUTED))
    lim_a = np.pi * (1 - limit / top)
    ax.text(1.12 * np.cos(lim_a), 1.12 * np.sin(lim_a), f"máx {limit * 100:.0f} %", ha="center", fontsize=9,
            color=hexc(OKGREEN), fontweight="bold")
    ax.set_xlim(-1.3, 1.3)
    ax.set_ylim(-0.25, 1.3)
    ax.set_aspect("equal")
    ax.axis("off")
    fig.savefig(path, transparent=True)
    plt.close(fig)


# ---------- Primitivas de diseño ----------
def _alpha(shape, pct):
    clr = shape.fill._xPr.find(".//a:solidFill/a:srgbClr", NS)
    if clr is not None:
        a = etree.SubElement(clr, "{%s}alpha" % NS["a"])
        a.set("val", str(int(pct * 1000)))


def _shadow(shape, blur=9, dist=3, alpha=18):
    spPr = shape._element.spPr
    for e in spPr.findall("a:effectLst", NS):
        spPr.remove(e)
    eff = etree.SubElement(spPr, "{%s}effectLst" % NS["a"])
    sh = etree.SubElement(eff, "{%s}outerShdw" % NS["a"], blurRad=str(blur * 12700), dist=str(dist * 12700),
                          dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(sh, "{%s}srgbClr" % NS["a"], val="0F1F3D")
    etree.SubElement(c, "{%s}alpha" % NS["a"], val=str(alpha * 1000))


def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, alpha=None, shadow=False):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if alpha is not None:
            _alpha(sh, alpha)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    if shadow:
        _shadow(sh)
    else:
        sh.shadow.inherit = False
    return sh


def rrect(s, x, y, w, h, fill, radius=0.08, **kw):
    return rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, **kw)


def card(s, x, y, w, h, fill=WHITE, radius=0.05, accent=None):
    c = rrect(s, x, y, w, h, fill, radius, line=LINE if fill == WHITE else None, shadow=True)
    if accent is not None:
        rect(s, x + Inches(0.18), y, Inches(0.6), Inches(0.06), accent)
    return c


def grad(sh, c1, c2, angle=90):
    sh.fill.gradient()
    sh.fill.gradient_angle = angle
    st = sh.fill.gradient_stops
    st[0].color.rgb, st[0].position = c1, 0
    st[1].color.rgb, st[1].position = c2, 1


def text(s, x, y, w, h, t, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
         spacing=None):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, ln in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        r.font.color.rgb, r.font.name = color, FONT
        if spacing:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def circle_num(s, x, y, d, label, fill=NAVY, color=WHITE, size=14):
    c = rect(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = str(label)
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), True, FONT, color
    return c


def pill(s, x, y, w, h, label, fill, color=WHITE, size=10):
    p = rrect(s, x, y, w, h, fill, 0.5)
    tf = p.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = tf.paragraphs[0]
    pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), True, FONT, color
    return p


def base(prs, tag, titulo, subtitulo, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # banda lateral y esquina decorativa
    grad(rect(s, 0, 0, Inches(0.22), H, NAVY), NAVY, BLUE, 90)
    rect(s, 0, 0, Inches(0.22), Inches(1.0), ACCENT)
    tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, W - Inches(1.6), 0, Inches(1.6), Inches(1.1))
    tri.rotation = 0
    tri.flip = None
    tri.fill.solid()
    tri.fill.fore_color.rgb = NAVY
    tri.line.fill.background()
    tri.shadow.inherit = False
    tri._element.spPr.find("a:xfrm", NS).set("flipH", "1")
    circle_num(s, W - Inches(0.95), Inches(0.14), Inches(0.62), f"{n:02d}", ACCENT, NAVY, 15)
    text(s, Inches(0.6), Inches(0.28), Inches(8), Inches(0.3), tag.upper(), 11, True, BLUE, spacing=200)
    text(s, Inches(0.6), Inches(0.52), Inches(11), Inches(0.6), titulo, 30, True, NAVY)
    text(s, Inches(0.6), Inches(1.08), Inches(11.5), Inches(0.35), subtitulo, 13, False, MUTED)
    rect(s, Inches(0.6), H - Inches(0.42), W - Inches(1.2), Emu(9525), G1)
    text(s, Inches(0.6), H - Inches(0.38), Inches(8), Inches(0.3), f"{AVALUADOR}  ·  {RAA}", 10, True, NAVY)
    text(s, W - Inches(6.6), H - Inches(0.38), Inches(6.0), Inches(0.3), f"Estudio de mercado rural  ·  {n} de {total}",
         10, False, MUTED, PP_ALIGN.RIGHT)
    return s


def kpi_card(s, x, y, w, h, icon, label, value, sub=None, fill=WHITE, vcolor=NAVY, size=24):
    card(s, x, y, w, h, fill, accent=ACCENT)
    circle_num(s, x + Inches(0.22), y + Inches(0.25), Inches(0.55), icon, NAVY if fill == WHITE else WHITE,
               WHITE if fill == WHITE else NAVY, 15)
    text(s, x + Inches(0.9), y + Inches(0.2), w - Inches(1.0), Inches(0.3), label, 10, True, MUTED, spacing=100)
    text(s, x + Inches(0.9), y + Inches(0.46), w - Inches(1.0), Inches(0.6), value, size, True, vcolor,
         anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(s, x + Inches(0.9), y + h - Inches(0.42), w - Inches(1.0), Inches(0.35), sub, 10, False, MUTED)


def progress(s, x, y, w, h, frac, fill=NAVY, track=G2, marker=None):
    rrect(s, x, y, w, h, track, 0.5)
    if frac > 0:
        rrect(s, x, y, max(int(w * min(frac, 1)), int(h)), h, fill, 0.5)
    if marker is not None:
        mx = x + int(w * marker)
        rect(s, mx - Emu(12000), y - Inches(0.05), Emu(24000), h + Inches(0.1), OKGREEN)


def tabla(s, x, y, w, rows, widths, hdr_size=10.5, size=10.5, row_h=0.36, align=None, hl_last=False, colors=None):
    gt = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    t = gt.table
    tot = sum(widths)
    for i, wd in enumerate(widths):
        t.columns[i].width = int(w * wd / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            last = hl_last and i == len(rows) - 1
            cell.fill.fore_color.rgb = NAVY if i == 0 else (GREEN if last else (WHITE if i % 2 else RGBColor(0xF7, 0xF9, 0xFC)))
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            p.alignment = PP_ALIGN.CENTER if i == 0 else (align[j] if align else PP_ALIGN.LEFT)
            r = p.add_run()
            r.text = str(v)
            r.font.name = FONT
            r.font.size = Pt(hdr_size if i == 0 else size)
            r.font.bold = i == 0 or last
            col = WHITE if i == 0 else INK
            if colors and i > 0 and (i, j) in colors:
                col = colors[(i, j)]
            r.font.color.rgb = col
    return gt


def estilo_chart(ch, legend=True):
    ch.font.size, ch.font.name = Pt(11), FONT
    ch.font.color.rgb = INK2
    ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False


def recortar(t, n):
    t = str(t or "")
    return t if len(t) <= n else t[: n - 1].rstrip() + "…"


def es_descartada(o):
    return o["comp"].upper().startswith("DESCARTADA")


# ---------- Diapositivas ----------
def construir(D, salida):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    TOTAL = 10
    of, st, sel, py, lq = D["ofertas"], D["est"], D["sel"], D["proy"], D["liq"]
    cumple = str(sel["cumple"]).upper().startswith("S")
    nvig = sum(1 for o in of if not es_descartada(o))
    tmp = Path(tempfile.mkdtemp())
    topo, gauge = tmp / "topo.png", tmp / "gauge.png"
    img_topo(topo)
    img_gauge(gauge, sel["cv"], D["cv_max"])

    # ===== 1. Portada =====
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, 0, 0, W, H, NAVY)
    grad(bg, DEEP, BLUE, 35)
    s.shapes.add_picture(str(topo), 0, 0, W, H)
    big = rect(s, Inches(8.4), Inches(-1.8), Inches(7.2), Inches(7.2), BLUE, shape=MSO_SHAPE.OVAL, alpha=22)
    rect(s, Inches(10.9), Inches(3.9), Inches(3.6), Inches(3.6), ACCENT, shape=MSO_SHAPE.OVAL, alpha=14)
    rect(s, Inches(0.8), Inches(1.0), Inches(0.9), Inches(0.07), ACCENT)
    text(s, Inches(0.8), Inches(1.15), Inches(8), Inches(0.4), "GESTIÓN VALUATORIA  ·  RES. 941", 13, True, ACCENT, spacing=300)
    text(s, Inches(0.8), Inches(1.7), Inches(10.5), Inches(2.3), ["Estudio de mercado", "predios rurales"], 54, True, WHITE)
    text(s, Inches(0.8), Inches(4.0), Inches(9.5), Inches(0.6), recortar(D["direccion"].split(". Coord")[0], 100), 20, False, G2)
    pill(s, Inches(0.8), Inches(4.75), Inches(3.9), Inches(0.42), f"{D['municipio'].split(' (')[0]}  ·  {D['depto']}", BLUE, WHITE, 12)
    gl = rrect(s, Inches(0.8), Inches(5.65), Inches(11.7), Inches(1.35), WHITE, 0.12, alpha=9, line=SKY)
    text(s, Inches(1.1), Inches(5.78), Inches(5), Inches(0.3), "AVALUADOR", 10, True, ACCENT, spacing=250)
    text(s, Inches(1.1), Inches(6.05), Inches(6.5), Inches(0.5), AVALUADOR, 25, True, WHITE)
    text(s, Inches(1.1), Inches(6.52), Inches(6), Inches(0.35), RAA, 14, False, G2)
    rect(s, Inches(7.9), Inches(5.85), Emu(12000), Inches(0.95), SKY)
    text(s, Inches(8.2), Inches(5.78), Inches(2), Inches(0.3), "FECHA", 10, True, ACCENT, spacing=250)
    text(s, Inches(8.2), Inches(6.05), Inches(2.4), Inches(0.5), D["fecha"], 24, True, WHITE)
    text(s, Inches(10.4), Inches(5.78), Inches(2), Inches(0.3), "SOLICITANTE", 10, True, ACCENT, spacing=250)
    text(s, Inches(10.4), Inches(6.12), Inches(2.1), Inches(0.5), recortar(D["solicitante"], 24), 14, True, WHITE)

    # ===== 2. Resumen ejecutivo =====
    s = base(prs, "Resumen ejecutivo", "El avalúo en una mirada", "Resultado del estudio y estado frente a la Resolución 941", 2, TOTAL)
    hero = rrect(s, Inches(0.6), Inches(1.6), Inches(4.6), Inches(5.05), NAVY, 0.05, shadow=True)
    grad(hero, DEEP, BLUE, 60)
    text(s, Inches(0.95), Inches(1.9), Inches(4), Inches(0.3), "VALOR TOTAL DEL AVALÚO", 11, True, ACCENT, spacing=250)
    text(s, Inches(0.95), Inches(2.35), Inches(4.0), Inches(1.0), mill(lq["total"], 1), 38, True, WHITE)
    text(s, Inches(0.95), Inches(3.35), Inches(4.0), Inches(0.4), cop(lq["total"]), 18, False, G2)
    rect(s, Inches(0.95), Inches(3.95), Inches(0.7), Emu(25000), ACCENT)
    text(s, Inches(0.95), Inches(4.15), Inches(3.9), Inches(1.2), lq["letras"].capitalize(), 13, False, G2, italic=True)
    pill(s, Inches(0.95), Inches(5.85), Inches(1.9), Inches(0.4),
         f"Terreno {num(lq['terreno'][2], 2)} ha" if lq["terreno"] else "Terreno", BLUE, WHITE, 11)
    pill(s, Inches(3.0), Inches(5.85), Inches(1.9), Inches(0.4), f"{len(D['constr'])} construcciones", BLUE, WHITE, 11)
    kw, kh = Inches(3.55), Inches(1.5)
    kpi_card(s, Inches(5.5), Inches(1.6), kw, kh, "$", "VALOR ADOPTADO", mill(fnum(st.get("VALOR ADOPTADO"))), "por hectárea", size=26)
    kpi_card(s, Inches(9.25), Inches(1.6), kw, kh, "#", "OFERTAS", f"{len(of)} · {nvig} vigente" + ("" if nvig == 1 else "s"),
             f"{len(of) - nvig} descartadas", size=20)
    kpi_card(s, Inches(5.5), Inches(3.3), kw, kh, "Σ", "COMBINACIÓN ELEGIDA", mill(sel["ha"]), f"{sel['nombre']} · {cop(sel['m2'])}/m²", size=26)
    kpi_card(s, Inches(9.25), Inches(3.3), kw, kh, "↔", "RANGO ORIENTATIVO", f"{num(py['baja'] / 1e6, 0)} – {num(py['alta'] / 1e6, 0)} M",
             f"amplitud {mill(py['amp'], 0)}", size=22)
    card(s, Inches(5.5), Inches(5.0), Inches(7.3), Inches(1.65), YELLOW if not cumple else GREEN)
    s.shapes.add_picture(str(gauge), Inches(5.6), Inches(5.05), Inches(2.7), Inches(1.5))
    text(s, Inches(8.4), Inches(5.12), Inches(4.3), Inches(0.3), "COEFICIENTE DE VARIACIÓN", 10, True, MUTED, spacing=100)
    text(s, Inches(8.4), Inches(5.38), Inches(4.3), Inches(0.6), pct(sel["cv"]), 30, True, OKGREEN if cumple else RED)
    text(s, Inches(8.4), Inches(5.98), Inches(4.3), Inches(0.65),
         (f"Cumple el máximo de {pct(D['cv_max'], 0)} (Res. 941)." if cumple else
          f"No cumple el máximo de {pct(D['cv_max'], 0)}: valores orientativos; se requieren más ofertas."), 11, False, INK2)

    # ===== 3. Predio y alcance =====
    s = base(prs, "Predio y método", "Predio objeto y alcance", "Identificación del inmueble y marco metodológico", 3, TOTAL)
    card(s, Inches(0.6), Inches(1.6), Inches(6.2), Inches(5.05))
    ident = [("DIRECCIÓN / UBICACIÓN", D["direccion"]), ("DEPARTAMENTO", D["depto"]), ("MUNICIPIO", D["municipio"]),
             ("SOLICITANTE", D["solicitante"]), ("REFERENCIA DE LIQUIDACIÓN", D["referencia"])]
    y = 1.8
    for k, v in ident:
        hh = 0.85 if k.startswith("DIRECCIÓN") else 0.55
        rect(s, Inches(0.85), Inches(y + 0.05), Inches(0.08), Inches(hh - 0.15), ACCENT)
        text(s, Inches(1.05), Inches(y), Inches(5.5), Inches(0.25), k, 9, True, MUTED, spacing=120)
        text(s, Inches(1.05), Inches(y + 0.23), Inches(5.6), Inches(hh - 0.2), recortar(str(v).replace("Referencia de liquidación:", "").strip(), 130), 13, True, NAVY)
        y += hh + 0.08
    if lq["terreno"]:
        g = rrect(s, Inches(0.85), Inches(5.6), Inches(2.6), Inches(0.9), NAVY, 0.15)
        text(s, Inches(1.0), Inches(5.63), Inches(2.4), Inches(0.25), "SUPERFICIE", 9, True, ACCENT, spacing=150)
        text(s, Inches(1.0), Inches(5.85), Inches(2.4), Inches(0.6), f"{num(lq['terreno'][2], 2)} ha", 26, True, WHITE)
        g2 = rrect(s, Inches(3.6), Inches(5.6), Inches(2.95), Inches(0.9), GREEN, 0.15)
        text(s, Inches(3.75), Inches(5.63), Inches(2.7), Inches(0.25), "FECHA DEL ESTUDIO", 9, True, MUTED, spacing=150)
        text(s, Inches(3.75), Inches(5.85), Inches(2.7), Inches(0.6), D["fecha"], 24, True, NAVY)
    mets = [("Método de mercado", "Comparación de ofertas depuradas por % de negociación."),
            ("Estadística", "Promedio, desviación, CV, límites y asimetría."),
            ("Res. 941 · rural", f"CV máximo {pct(D['cv_max'], 0)}; sin homogenización por factores."),
            ("Combinaciones", "Grupos de 3 o más ofertas con el menor CV."),
            ("Construcciones", "Depreciación Ross-Heidecke por estado."),
            ("Liquidación", "Terreno + construcciones, redondeo a la centena.")]
    for i, (k, v) in enumerate(mets):
        cx, cy = Inches(7.05 + (i % 2) * 2.95), Inches(1.6 + (i // 2) * 1.72)
        card(s, cx, cy, Inches(2.8), Inches(1.57))
        circle_num(s, cx + Inches(0.18), cy + Inches(0.2), Inches(0.5), i + 1, NAVY if i % 2 == 0 else BLUE, WHITE, 15)
        text(s, cx + Inches(0.8), cy + Inches(0.22), Inches(1.9), Inches(0.5), k, 13, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + Inches(0.18), cy + Inches(0.85), Inches(2.5), Inches(0.7), v, 10.5, False, INK2)

    # ===== 4. Ofertas =====
    s = base(prs, "Mercado", "Ofertas de mercado", "Hoja MERCADO RURAL · valores depurados por % de negociación", 4, TOTAL)
    card(s, Inches(0.6), Inches(1.55), Inches(12.15), Inches(5.2))
    rows = [["No.", "Ubicación", "Área ha", "Valor pedido", "% neg.", "Valor depurado", "Valor $/ha", "Comparabilidad"]]
    cols = {}
    for i, o in enumerate(of, start=1):
        etiqueta = "ALTA" if not es_descartada(o) else "DESCARTADA"
        rows.append([o["no"], recortar(o["ubic"], 52), num(o["area"], 2), mill(o["pedido"], 0), pct(o["neg"], 0),
                     mill(o["depurado"], 0), mill(o["ha"], 2), etiqueta if es_descartada(o) else o["comp"].split(" ")[0]])
        cols[(i, 7)] = MUTED if es_descartada(o) else OKGREEN
        if not es_descartada(o):
            cols[(i, 6)] = OKGREEN
    n = len(of)
    tabla(s, Inches(0.75), Inches(1.7), Inches(11.85), rows, [0.5, 4.3, 0.9, 1.2, 0.7, 1.3, 1.3, 1.5],
          size=10.5, row_h=min(0.5, 4.6 / (n + 1)), colors=cols,
          align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * 5 + [PP_ALIGN.CENTER])
    text(s, Inches(0.75), Inches(6.35), Inches(11.8), Inches(0.35),
         "M = millones de pesos. ALTA = comparable en uso agropecuario · DESCARTADA = turístico, playa, campestre o dato no confiable.",
         10, False, MUTED, italic=True)

    # ===== 5. Gráfico y estadísticos =====
    s = base(prs, "Estadística", "Valor por hectárea y estadísticos", "Muestra general · $ millones por ha (escala logarítmica)", 5, TOTAL)
    card(s, Inches(0.6), Inches(1.55), Inches(7.9), Inches(5.2))
    cd = CategoryChartData()
    cd.categories = [f"Of. {o['no']}" for o in of]
    cd.add_series("Valor $/ha (millones)", [round(o["ha"] / 1e6, 3) for o in of])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.65), Inches(7.7), Inches(4.7), cd)
    ch = gf.chart
    estilo_chart(ch, False)
    ch.plots[0].gap_width = 45
    ser = ch.plots[0].series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = NAVY
    for i, o in enumerate(of):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = G1 if es_descartada(o) else NAVY
    ser.data_labels.show_value = True
    ser.data_labels.font.size = Pt(10)
    ser.data_labels.font.bold = True
    ser.data_labels.number_format = '#,##0.0'
    ser.data_labels.number_format_is_linked = False
    ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.minimum_scale = 1
    scaling = va._element.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling')
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    lb = etree.Element(f'{{{ns}}}logBase')
    lb.set('val', '10')
    scaling.insert(0, lb)
    if scaling.find(f'{{{ns}}}orientation') is None:
        ori = etree.Element(f'{{{ns}}}orientation')
        ori.set('val', 'minMax')
        scaling.insert(1, ori)
    ch.category_axis.format.line.color.rgb = G1
    rect(s, Inches(0.95), Inches(6.42), Inches(0.18), Inches(0.18), NAVY)
    text(s, Inches(1.18), Inches(6.36), Inches(2.5), Inches(0.3), "Oferta vigente", 10, False, INK2)
    rect(s, Inches(3.0), Inches(6.42), Inches(0.18), Inches(0.18), G1)
    text(s, Inches(3.23), Inches(6.36), Inches(2.5), Inches(0.3), "Oferta descartada", 10, False, INK2)
    items = [("PROMEDIO", lambda v: mill(v)), ("DESVIACIÓN ESTÁNDAR", lambda v: mill(v)),
             ("COEFICIENTE DE VARIACIÓN", pct), ("COEFICIENTE DE ASIMETRÍA", lambda v: num(v, 2)),
             ("LÍMITE INFERIOR", lambda v: mill(v)), ("LÍMITE SUPERIOR", lambda v: mill(v))]
    for i, (k, f) in enumerate(items):
        cx, cy = Inches(8.75 + (i % 2) * 2.05), Inches(1.55 + (i // 2) * 1.12)
        v = st.get(k)
        bad = k == "COEFICIENTE DE VARIACIÓN" and fnum(v) > D["cv_max"]
        card(s, cx, cy, Inches(1.95), Inches(1.0), accent=RED if bad else ACCENT)
        text(s, cx + Inches(0.1), cy + Inches(0.12), Inches(1.8), Inches(0.35), k.capitalize(), 9, True, MUTED)
        text(s, cx + Inches(0.1), cy + Inches(0.42), Inches(1.8), Inches(0.5), f(fnum(v)) if v is not None else "—", 17, True,
             RED if bad else NAVY, anchor=MSO_ANCHOR.MIDDLE)
    hv = rrect(s, Inches(8.75), Inches(4.95), Inches(4.0), Inches(1.8), NAVY, 0.07, shadow=True)
    grad(hv, DEEP, BLUE, 45)
    text(s, Inches(9.0), Inches(5.1), Inches(3.6), Inches(0.3), "VALOR ADOPTADO", 10, True, ACCENT, spacing=250)
    text(s, Inches(9.0), Inches(5.45), Inches(3.6), Inches(0.8), mill(fnum(st.get("VALOR ADOPTADO"))), 36, True, WHITE)
    text(s, Inches(9.0), Inches(6.2), Inches(3.6), Inches(0.4), "por hectárea · muestra general", 11, False, G2)

    # ===== 6. Combinaciones =====
    s = base(prs, "Combinaciones", "Combinaciones de ofertas", f"Grupos de 3 o más ofertas · CV máximo admisible {pct(D['cv_max'], 0)} (Res. 941)", 6, TOTAL)
    nc = max(len(D["combos"]), 1)
    cw = (12.15 - 0.25 * (nc - 1)) / nc
    topcv = max([c_["cv"] for c_ in D["combos"]] + [D["cv_max"] * 2]) * 1.15
    for i, c_ in enumerate(D["combos"]):
        cx = Inches(0.6 + i * (cw + 0.25))
        es_sel = c_["nombre"] == sel["nombre"]
        card(s, cx, Inches(1.6), Inches(cw), Inches(3.35), accent=ACCENT)
        if es_sel:
            rrect(s, cx, Inches(1.6), Inches(cw), Inches(3.35), None, 0.05, line=BLUE)
            pill(s, cx + Inches(cw - 1.6), Inches(1.78), Inches(1.45), Inches(0.3), "SELECCIONADA", BLUE, WHITE, 9)
        text(s, cx + Inches(0.2), Inches(1.75), Inches(cw - 1.8), Inches(0.4), c_["nombre"], 16, True, NAVY)
        text(s, cx + Inches(0.2), Inches(2.2), Inches(cw - 0.4), Inches(0.3), "OFERTAS: " + str(c_["ofertas"]), 10, True, MUTED, spacing=100)
        ok = str(c_["cumple"]).upper().startswith("S")
        text(s, cx + Inches(0.2), Inches(2.55), Inches(2.4), Inches(0.9), pct(c_["cv"]), 36, True, OKGREEN if ok else RED)
        pill(s, cx + Inches(cw - 1.4), Inches(2.85), Inches(1.2), Inches(0.34), "CUMPLE" if ok else "NO CUMPLE", OKGREEN if ok else RED, WHITE, 9)
        progress(s, cx + Inches(0.2), Inches(3.62), Inches(cw - 0.4), Inches(0.16), c_["cv"] / topcv,
                 OKGREEN if ok else RED, G2, D["cv_max"] / topcv)
        text(s, cx + Inches(0.2), Inches(3.8), Inches(cw - 0.4), Inches(0.28), f"│ máximo admisible {pct(D['cv_max'], 0)}", 9, False, OKGREEN)
        text(s, cx + Inches(0.2), Inches(4.2), Inches(1.8), Inches(0.28), "Promedio $/ha", 9, True, MUTED)
        text(s, cx + Inches(0.2), Inches(4.43), Inches(1.9), Inches(0.4), mill(c_["prom"]), 15, True, NAVY)
        text(s, cx + Inches(cw / 2), Inches(4.2), Inches(1.8), Inches(0.28), "Desv. estándar", 9, True, MUTED)
        text(s, cx + Inches(cw / 2), Inches(4.43), Inches(1.9), Inches(0.4), mill(c_["de"]), 15, True, NAVY)
    band = rrect(s, Inches(0.6), Inches(5.2), Inches(12.15), Inches(1.5), NAVY, 0.08, shadow=True)
    grad(band, DEEP, BLUE, 0)
    text(s, Inches(0.9), Inches(5.35), Inches(4), Inches(0.3), "COMBINACIÓN SELECCIONADA", 10, True, ACCENT, spacing=250)
    text(s, Inches(0.9), Inches(5.7), Inches(4.2), Inches(0.5), sel["nombre"], 24, True, WHITE)
    text(s, Inches(0.9), Inches(6.2), Inches(4.3), Inches(0.4), f"Ofertas {sel['ofertas']}", 12, False, G2)
    for i, (k, v, col) in enumerate([("CV", pct(sel["cv"]), RED if not cumple else ACCENT), ("¿CUMPLE?", sel["cumple"], RED if not cumple else ACCENT),
                                    ("VALOR POR HA", mill(sel["ha"]), WHITE), ("VALOR POR M²", cop(sel["m2"]), WHITE)]):
        x = Inches(5.3 + i * 1.95)
        rect(s, x - Inches(0.12), Inches(5.45), Emu(12000), Inches(1.0), SKY, alpha=60)
        text(s, x, Inches(5.4), Inches(1.8), Inches(0.3), k, 9, True, SKY, spacing=150)
        text(s, x, Inches(5.75), Inches(1.85), Inches(0.7), v, 20, True, col)

    # ===== 7. Proyección =====
    s = base(prs, "Proyección", "Proyección del valor esperado", "Rango de orientación entre dos combinaciones de referencia · sin valor puntual", 7, TOTAL)
    kpi_card(s, Inches(0.6), Inches(1.6), Inches(3.9), Inches(1.45), "↓", "REFERENCIA BAJA · $/ha", mill(py["baja"]), str(py["nbaja"]), size=26)
    kpi_card(s, Inches(4.7), Inches(1.6), Inches(3.9), Inches(1.45), "↑", "REFERENCIA ALTA · $/ha", mill(py["alta"]), str(py["nalta"]), size=26)
    kpi_card(s, Inches(8.8), Inches(1.6), Inches(3.95), Inches(1.45), "↔", "AMPLITUD DEL RANGO", mill(py["amp"]),
             f"{pct(py['amp'] / py['baja'], 0)} de la referencia baja" if py["baja"] else "", size=26)
    card(s, Inches(0.6), Inches(3.25), Inches(12.15), Inches(2.0))
    text(s, Inches(0.85), Inches(3.35), Inches(8), Inches(0.3), "POSICIÓN DEL VALOR ADOPTADO FRENTE AL RANGO · $ millones por ha", 10, True, MUTED, spacing=100)
    lo, hi = py["baja"] / 1e6, py["alta"] / 1e6
    adop = fnum(st.get("VALOR ADOPTADO")) / 1e6
    vmin, vmax = min(lo, adop) * 0.8, max(hi, adop) * 1.2
    bx, bw, by = Inches(1.0), Inches(11.4), Inches(4.3)
    px = lambda v: int(bx + bw * (v - vmin) / (vmax - vmin))
    rrect(s, bx, by, bw, Inches(0.36), G2, 0.5)
    seg = rrect(s, px(lo), by, max(px(hi) - px(lo), Inches(0.1)), Inches(0.36), NAVY, 0.5)
    grad(seg, NAVY, BLUE, 0)
    for v, lab in ((lo, f"Baja {num(lo, 1)}"), (hi, f"Alta {num(hi, 1)}")):
        circle_num(s, px(v) - Inches(0.14), by + Inches(0.04), Inches(0.28), "", WHITE, NAVY)
        text(s, px(v) - Inches(0.9), by + Inches(0.45), Inches(1.8), Inches(0.35), lab, 12, True, NAVY, PP_ALIGN.CENTER)
    rect(s, px(adop) - Emu(15000), by - Inches(0.35), Emu(30000), Inches(0.35), OKGREEN)
    circle_num(s, px(adop) - Inches(0.17), by + Inches(0.01), Inches(0.34), "", ACCENT, NAVY)
    text(s, px(adop) - Inches(1.3), by - Inches(0.7), Inches(2.6), Inches(0.35), f"Adoptado {num(adop, 1)}", 12, True, OKGREEN, PP_ALIGN.CENTER)
    rows = [["Referencia", "Combinación", "Ofertas", "Promedio $/ha", "$/m²", "CV", "¿Cumple?"]]
    cols = {}
    for i, r_ in enumerate(py["ref"], start=1):
        rows.append([r_[0], r_[1], r_[2], mill(fnum(r_[3])), cop(fnum(r_[4])), pct(fnum(r_[5])), r_[6]])
        cols[(i, 6)] = RED if str(r_[6]).upper() == "NO" else OKGREEN
    tabla(s, Inches(0.6), Inches(5.45), Inches(12.15), rows, [1.6, 1.6, 1.3, 1.6, 1.2, 0.9, 1.0], row_h=0.34, colors=cols,
          align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER])
    text(s, Inches(0.6), Inches(6.55), Inches(12.15), Inches(0.5), recortar(py["nota"], 260), 10, False, MUTED, italic=True)

    # ===== 8. Depreciación =====
    s = base(prs, "Construcciones", "Depreciación de construcciones", "Método Ross-Heidecke · barra = valor remanente (1 − depreciación)", 8, TOTAL)
    nco = len(D["constr"])
    rh = min(0.6, 4.3 / (nco + 1))
    card(s, Inches(0.6), Inches(1.55), Inches(12.15), Inches(0.5 + rh * (nco + 1) + 0.2))
    heads = [("ÍTEM", 0.8, 2.3), ("EDAD / VIDA", 3.15, 1.2), ("ESTADO", 4.35, 0.8), ("REPOSICIÓN $/m²", 5.2, 1.5),
             ("VALOR REMANENTE", 6.95, 2.6), ("ÁREA m²", 9.7, 0.9), ("VALOR TOTAL", 10.6, 2.0)]
    for t_, x_, w_ in heads:
        text(s, Inches(x_), Inches(1.65), Inches(w_), Inches(0.3), t_, 9, True, MUTED, PP_ALIGN.RIGHT if t_ in ("VALOR TOTAL", "ÁREA m²") else PP_ALIGN.LEFT, spacing=100)
    y = 2.05
    for i, c_ in enumerate(D["constr"]):
        if i % 2 == 0:
            rect(s, Inches(0.75), Inches(y), Inches(11.85), Inches(rh), RGBColor(0xF7, 0xF9, 0xFC))
        mid = Inches(y)
        text(s, Inches(0.8), mid, Inches(2.3), Inches(rh), c_["item"], 12, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(3.15), mid, Inches(1.2), Inches(rh), f"{c_['edad']} / {c_['vida']} años", 11, False, INK2, anchor=MSO_ANCHOR.MIDDLE)
        pill(s, Inches(4.38), Inches(y + rh / 2 - 0.14), Inches(0.5), Inches(0.28), str(c_["estado"]), BLUE if c_["estado"] <= 3 else AMBER, WHITE, 10)
        text(s, Inches(5.2), mid, Inches(1.5), Inches(rh), cop(c_["repos"]), 11, False, INK2, anchor=MSO_ANCHOR.MIDDLE)
        progress(s, Inches(6.95), Inches(y + rh / 2 - 0.08), Inches(1.9), Inches(0.16), c_["fd"], ACCENT if c_["fd"] >= 0.5 else AMBER)
        text(s, Inches(8.9), mid, Inches(0.7), Inches(rh), pct(c_["fd"], 0), 11, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(9.7), mid, Inches(0.9), Inches(rh), num(c_["area"], 0), 11, False, INK2, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
        text(s, Inches(10.6), mid, Inches(1.95), Inches(rh), cop(c_["total"]), 12, True, NAVY, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
        y += rh
    tot = rrect(s, Inches(0.75), Inches(y + 0.05), Inches(11.85), Inches(rh), NAVY, 0.2)
    grad(tot, DEEP, BLUE, 0)
    text(s, Inches(0.95), Inches(y + 0.05), Inches(4), Inches(rh), "TOTAL CONSTRUCCIONES", 12, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE, spacing=100)
    text(s, Inches(9.7), Inches(y + 0.05), Inches(0.9), Inches(rh), num(sum(c_["area"] for c_ in D["constr"]), 0), 12, True, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    text(s, Inches(10.6), Inches(y + 0.05), Inches(1.95), Inches(rh), cop(D["constr_total"]), 14, True, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)

    # ===== 9. Liquidación =====
    s = base(prs, "Liquidación", "Liquidación del avalúo", "Terreno + construcciones · totales redondeados a la centena", 9, TOTAL)
    ban = rrect(s, Inches(0.6), Inches(1.6), Inches(12.15), Inches(1.35), NAVY, 0.08, shadow=True)
    grad(ban, DEEP, BLUE, 0)
    text(s, Inches(0.9), Inches(1.72), Inches(5), Inches(0.3), "VALOR TOTAL DEL AVALÚO", 10, True, ACCENT, spacing=250)
    text(s, Inches(0.9), Inches(2.0), Inches(6), Inches(0.8), cop(lq["total"]), 38, True, WHITE)
    text(s, Inches(7.0), Inches(1.78), Inches(5.5), Inches(0.3), "SON", 10, True, ACCENT, spacing=250)
    text(s, Inches(7.0), Inches(2.05), Inches(5.6), Inches(0.8), lq["letras"], 13, True, WHITE)
    rows = [["Concepto", "Unidad", "Cantidad", "Valor unitario", "Valor total"]]
    if lq["terreno"]:
        t_ = lq["terreno"]
        rows.append(["Terreno", t_[1], num(t_[2], 2), cop(t_[3]), cop(t_[4])])
    for t_ in lq["principal"] + lq["anexas"]:
        rows.append([t_[0].title(), t_[1], num(t_[2], 0), cop(t_[3]), cop(t_[4])])
    rows.append(["VALOR TOTAL DEL AVALÚO", "", "", "", cop(lq["total"])])
    tabla(s, Inches(0.6), Inches(3.15), Inches(7.6), rows, [3.0, 0.8, 1.0, 1.6, 1.8], size=10.5,
          row_h=min(0.42, 3.55 / len(rows)), hl_last=True,
          align=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT])
    terr = lq["terreno"][4] if lq["terreno"] else 0
    partes = [("Terreno", terr, NAVY), ("Constr. principal", lq["tot_p"], BLUE), ("Constr. anexas", lq["tot_a"], SKY)]
    card(s, Inches(8.4), Inches(3.15), Inches(4.35), Inches(3.55))
    text(s, Inches(8.6), Inches(3.25), Inches(4), Inches(0.3), "COMPOSICIÓN DEL VALOR", 10, True, MUTED, spacing=100)
    cd = CategoryChartData()
    cd.categories = [p_[0] for p_ in partes]
    cd.add_series("Valor", [p_[1] / 1e6 for p_ in partes])
    gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(8.45), Inches(3.5), Inches(2.4), Inches(2.4), cd)
    ch = gf.chart
    estilo_chart(ch, False)
    ch.plots[0].has_data_labels = False
    for i, (_, _, col) in enumerate(partes):
        pt = ch.plots[0].series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = col
    for i, (nm, v, col) in enumerate(partes):
        yy = Inches(3.95 + i * 0.7)
        rect(s, Inches(10.95), yy + Inches(0.08), Inches(0.16), Inches(0.16), col)
        text(s, Inches(11.18), yy - Inches(0.03), Inches(1.6), Inches(0.3), nm, 10, True, INK2)
        text(s, Inches(11.18), yy + Inches(0.22), Inches(1.6), Inches(0.3),
             pct(v / lq["total"]) if lq["total"] else "—", 14, True, NAVY)
    text(s, Inches(8.6), Inches(6.05), Inches(4.0), Inches(0.5), "Construcciones: " + mill(lq["tot_p"] + lq["tot_a"], 1), 11, False, MUTED)

    # ===== 10. Observaciones =====
    s = base(prs, "Conclusiones", "Observaciones y salvedades", "Alcance y limitaciones del estudio · responsable técnico", 10, TOTAL)
    obs = [x.strip() for x in D["obs"].replace("OBSERVACIONES:", "").split("\n") if x.strip()][:5]
    y = 1.6
    hh = min(1.0, 5.1 / max(len(obs), 1) - 0.1)
    for i, o_ in enumerate(obs, start=1):
        o_ = o_.lstrip("0123456789) ").strip()
        card(s, Inches(0.6), Inches(y), Inches(8.2), Inches(hh))
        circle_num(s, Inches(0.78), Inches(y + hh / 2 - 0.22), Inches(0.44), i, NAVY, WHITE, 13)
        text(s, Inches(1.35), Inches(y), Inches(7.35), Inches(hh), recortar(o_, 430), 10.5, False, INK2, anchor=MSO_ANCHOR.MIDDLE)
        y += hh + 0.1
    sg = rrect(s, Inches(9.05), Inches(1.6), Inches(3.7), Inches(5.1), NAVY, 0.06, shadow=True)
    grad(sg, DEEP, BLUE, 60)
    text(s, Inches(9.3), Inches(1.95), Inches(3.2), Inches(0.3), "AVALUADOR", 10, True, ACCENT, PP_ALIGN.CENTER, spacing=300)
    circle_num(s, Inches(10.4), Inches(2.4), Inches(1.0), "".join(w_[0] for w_ in AVALUADOR.split()[:2]), BLUE, WHITE, 24)
    text(s, Inches(9.3), Inches(3.55), Inches(3.2), Inches(0.9), AVALUADOR, 20, True, WHITE, PP_ALIGN.CENTER)
    text(s, Inches(9.3), Inches(4.5), Inches(3.2), Inches(0.4), RAA, 13, True, ACCENT, PP_ALIGN.CENTER)
    rect(s, Inches(9.5), Inches(5.75), Inches(2.8), Emu(12000), SKY)
    text(s, Inches(9.3), Inches(5.8), Inches(3.2), Inches(0.3), "Firma", 10, False, SKY, PP_ALIGN.CENTER)
    text(s, Inches(9.3), Inches(6.2), Inches(3.2), Inches(0.35), "Fecha del estudio: " + D["fecha"], 11, True, WHITE, PP_ALIGN.CENTER)

    prs.save(salida)
    return salida


if __name__ == "__main__":
    xlsx = Path(sys.argv[1] if len(sys.argv) > 1 else "Estudio_Mercado_Rural_Puente_Bomba_Tigreras.xlsx")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "Presentacion_Estudio_Mercado_Rural.pptx")
    data = leer(xlsx)
    if not data["ofertas"] or data["liq"]["total"] == 0:
        sys.exit("El libro no tiene valores calculados: ábralo en Excel, guárdelo y vuelva a ejecutar.")
    print("Generada:", construir(data, out))
